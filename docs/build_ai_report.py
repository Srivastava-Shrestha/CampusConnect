"""
Build a styled Word document from docs/AI_ARCHITECTURE_DECISIONS.md.

Everything stays inside this repository - the input Markdown, this script, and
the generated .docx all live in MAY2026-Team-003/docs/.

Run it from the docs/ folder:

    cd MAY2026-Team-003/docs
    python build_ai_report.py

Then open the generated .docx in Word and use File > Save As > PDF to produce
the final PDF. That is the same path we used for the Milestone 2 deliverable,
so the two documents come out looking like a matched set.

Styling matches the Milestone 2 house style:
  - navy table headers with white bold text
  - light blue banded rows
  - full grid borders on every table
  - Mermaid diagram sources rendered as labelled monospace boxes
"""

import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


# ----------------------------------------------------------------------------
# Configuration
# ----------------------------------------------------------------------------

HERE = os.path.dirname(os.path.abspath(__file__))
SOURCE_MD = os.path.join(HERE, "AI_ARCHITECTURE_DECISIONS.md")
OUTPUT_DOCX = os.path.join(HERE, "AI-Architecture-Decisions-Team-003-Nexmind.docx")

# Milestone 2 house style colours.
NAVY = RGBColor(0x1F, 0x38, 0x64)
HEADER_FILL = "1F3864"
BAND_FILL = "EEF2F8"
CODE_FILL = "F5F5F5"
MERMAID_FILL = "EEF2F8"

DOC_TITLE = "Campus Connect - AI Architecture Decisions"
DOC_SUBTITLE = "Team NexMind (Team-003) - IITM BS Software Engineering, May 2026"


# ----------------------------------------------------------------------------
# Low level docx helpers
# ----------------------------------------------------------------------------

def set_cell_shading(cell, hex_color):
    """Fill a table cell with a solid background colour."""
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shading)


def set_paragraph_shading(paragraph, hex_color):
    """Fill a whole paragraph with a solid background colour."""
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), hex_color)
    paragraph._p.get_or_add_pPr().append(shading)


def apply_grid_and_header(table, band=True):
    """Apply the navy header + banded body + full grid look to a table."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        line = OxmlElement("w:" + edge)
        line.set(qn("w:val"), "single")
        line.set(qn("w:sz"), "4")
        line.set(qn("w:color"), "B4C2DA")
        borders.append(line)
    table._tbl.tblPr.append(borders)

    # Header row: navy fill, white bold text.
    for cell in table.rows[0].cells:
        set_cell_shading(cell, HEADER_FILL)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.size = Pt(9)

    # Body rows: alternate banding, slightly smaller text.
    for index, row in enumerate(table.rows[1:], start=1):
        for cell in row.cells:
            if band and index % 2 == 0:
                set_cell_shading(cell, BAND_FILL)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)


# ----------------------------------------------------------------------------
# Inline Markdown -> Word runs
# ----------------------------------------------------------------------------

# Split on bold, inline code, and links, keeping the delimiters.
INLINE_PATTERN = re.compile(
    r"(\*\*[^*]+\*\*)"      # **bold**
    r"|(`[^`]+`)"           # `code`
    r"|(\[[^\]]+\]\([^)]+\))"  # [text](url)
)


def clean_text(text):
    """Normalise characters that do not belong in the Word deliverable."""
    text = text.replace("—", " - ")   # em dash
    text = text.replace("–", "-")     # en dash
    text = text.replace("‘", "'")
    text = text.replace("’", "'")
    text = text.replace("“", '"')
    text = text.replace("”", '"')
    text = text.replace("→", "->")
    text = text.replace("≥", ">=")
    text = text.replace("≤", "<=")
    text = text.replace("×", "x")
    return text


def add_inline_runs(paragraph, text):
    """Write text into a paragraph, honouring bold, inline code, and links."""
    text = clean_text(text)
    position = 0

    for match in INLINE_PATTERN.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position:match.start()])

        bold_part, code_part, link_part = match.groups()

        if bold_part is not None:
            run = paragraph.add_run(bold_part[2:-2])
            run.bold = True
        elif code_part is not None:
            run = paragraph.add_run(code_part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0xB0, 0x30, 0x60)
        elif link_part is not None:
            label = link_part[1:link_part.index("]")]
            run = paragraph.add_run(label)
            run.font.color.rgb = NAVY
            run.underline = True

        position = match.end()

    if position < len(text):
        paragraph.add_run(text[position:])


# ----------------------------------------------------------------------------
# Block level renderers
# ----------------------------------------------------------------------------

def add_heading(doc, text, level):
    """Add a navy heading at the requested level."""
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
    paragraph.paragraph_format.space_after = Pt(6)

    sizes = {1: 20, 2: 15, 3: 12, 4: 11}
    run = paragraph.add_run(clean_text(text))
    run.bold = True
    run.font.size = Pt(sizes.get(level, 11))
    run.font.color.rgb = NAVY
    return paragraph


def add_body_paragraph(doc, text):
    """Add a normal body paragraph."""
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    add_inline_runs(paragraph, text)
    return paragraph


def add_bullet(doc, text, ordered=False):
    """Add a bulleted or numbered list item."""
    style = "List Number" if ordered else "List Bullet"
    try:
        paragraph = doc.add_paragraph(style=style)
    except KeyError:
        paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    add_inline_runs(paragraph, text)
    return paragraph


def add_quote(doc, text):
    """Add an indented, italic block quote with a navy left feel."""
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Pt(24)
    paragraph.paragraph_format.space_after = Pt(6)
    add_inline_runs(paragraph, text)
    for run in paragraph.runs:
        run.italic = True
        run.font.color.rgb = NAVY
    return paragraph


def add_code_block(doc, lines, language):
    """Add a monospace block. Mermaid gets a labelled caption above it."""
    if language == "mermaid":
        caption = doc.add_paragraph()
        caption.paragraph_format.space_before = Pt(8)
        caption.paragraph_format.space_after = Pt(0)
        run = caption.add_run("Diagram (Mermaid source - paste into mermaid.live to render)")
        run.bold = True
        run.italic = True
        run.font.size = Pt(8)
        run.font.color.rgb = NAVY
        fill = MERMAID_FILL
    else:
        fill = CODE_FILL

    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Pt(12)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(10)
    set_paragraph_shading(paragraph, fill)

    for index, line in enumerate(lines):
        if index > 0:
            paragraph.add_run("\n")
        run = paragraph.add_run(line.replace("\t", "    "))
        run.font.name = "Consolas"
        run.font.size = Pt(8)
    return paragraph


def split_table_row(line):
    """Split a Markdown table row into its cell strings."""
    stripped = line.strip()
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|"):
        stripped = stripped[:-1]
    return [cell.strip() for cell in stripped.split("|")]


def is_table_separator(line):
    """True for the ---|---|--- row under a Markdown table header."""
    stripped = line.strip()
    if "|" not in stripped or "-" not in stripped:
        return False
    return re.fullmatch(r"[\s|:-]+", stripped) is not None


def add_table(doc, rows):
    """Render a list of Markdown rows as a styled Word table."""
    header = split_table_row(rows[0])
    body = [split_table_row(row) for row in rows[1:]]

    column_count = len(header)
    table = doc.add_table(rows=1, cols=column_count)

    for index, cell_text in enumerate(header):
        cell = table.rows[0].cells[index]
        cell.text = ""
        add_inline_runs(cell.paragraphs[0], cell_text)

    for body_row in body:
        cells = table.add_row().cells
        for index in range(column_count):
            value = body_row[index] if index < len(body_row) else ""
            cells[index].text = ""
            add_inline_runs(cells[index].paragraphs[0], value)

    apply_grid_and_header(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


# ----------------------------------------------------------------------------
# Markdown walker
# ----------------------------------------------------------------------------

def render_markdown(doc, markdown_text):
    """Walk the Markdown line by line and emit the matching Word blocks."""
    lines = markdown_text.split("\n")
    index = 0

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        # Blank line.
        if not stripped:
            index += 1
            continue

        # Horizontal rule.
        if re.fullmatch(r"-{3,}|\*{3,}|_{3,}", stripped):
            doc.add_paragraph()
            index += 1
            continue

        # Fenced code block.
        if stripped.startswith("```"):
            language = stripped[3:].strip().lower()
            block = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            index += 1
            add_code_block(doc, block, language)
            continue

        # Heading.
        heading_match = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading_match:
            level = len(heading_match.group(1))
            add_heading(doc, heading_match.group(2), level)
            index += 1
            continue

        # Table: a pipe row followed by a separator row.
        if "|" in stripped and index + 1 < len(lines) and is_table_separator(lines[index + 1]):
            table_rows = [lines[index]]
            index += 2
            while index < len(lines) and "|" in lines[index] and lines[index].strip():
                table_rows.append(lines[index])
                index += 1
            add_table(doc, table_rows)
            continue

        # Block quote.
        if stripped.startswith(">"):
            quote_lines = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote_lines.append(lines[index].strip().lstrip(">").strip())
                index += 1
            add_quote(doc, " ".join(part for part in quote_lines if part))
            continue

        # Unordered list item.
        bullet_match = re.match(r"^\s*[-*+]\s+(.*)$", line)
        if bullet_match:
            add_bullet(doc, bullet_match.group(1), ordered=False)
            index += 1
            continue

        # Ordered list item.
        number_match = re.match(r"^\s*\d+[.)]\s+(.*)$", line)
        if number_match:
            add_bullet(doc, number_match.group(1), ordered=True)
            index += 1
            continue

        # Plain paragraph: join wrapped lines until a blank or a new block.
        paragraph_lines = []
        while index < len(lines):
            current = lines[index]
            current_stripped = current.strip()
            if not current_stripped:
                break
            if current_stripped.startswith(("#", ">", "```")):
                break
            if re.match(r"^\s*[-*+]\s+", current) or re.match(r"^\s*\d+[.)]\s+", current):
                break
            if "|" in current_stripped and index + 1 < len(lines) and is_table_separator(lines[index + 1]):
                break
            paragraph_lines.append(current_stripped)
            index += 1
        add_body_paragraph(doc, " ".join(paragraph_lines))


# ----------------------------------------------------------------------------
# Document shell
# ----------------------------------------------------------------------------

def add_title_page(doc):
    """Add the cover block at the top of the document."""
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(40)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(DOC_TITLE)
    title_run.bold = True
    title_run.font.size = Pt(26)
    title_run.font.color.rgb = NAVY

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(30)
    subtitle_run = subtitle.add_run(DOC_SUBTITLE)
    subtitle_run.italic = True
    subtitle_run.font.size = Pt(11)
    subtitle_run.font.color.rgb = NAVY

    doc.add_page_break()


def set_base_font(doc):
    """Set the default body font for the whole document."""
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)


def main():
    if not os.path.exists(SOURCE_MD):
        raise SystemExit("Could not find " + SOURCE_MD)

    with open(SOURCE_MD, "r", encoding="utf-8") as handle:
        markdown_text = handle.read()

    # Drop the in-document table of contents; Word users generate their own,
    # and the anchor links do not survive the conversion.
    markdown_text = re.sub(
        r"## Table of contents.*?(?=\n---\n)",
        "",
        markdown_text,
        flags=re.DOTALL,
    )

    doc = Document()
    set_base_font(doc)
    add_title_page(doc)
    render_markdown(doc, markdown_text)
    doc.save(OUTPUT_DOCX)

    print("Wrote " + OUTPUT_DOCX)
    print("Open it in Word and use File > Save As > PDF for the final deliverable.")


if __name__ == "__main__":
    main()
